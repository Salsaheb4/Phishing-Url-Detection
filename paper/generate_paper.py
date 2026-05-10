"""Generate paper/URL-Project.docx and paper/URL-Project.pdf in IEEE conference format."""
from __future__ import annotations
import re
from pathlib import Path

ROOT  = Path(__file__).resolve().parent.parent
FIGS  = ROOT / "outputs" / "figures"
PAPER = Path(__file__).resolve().parent
DOCX  = PAPER / "URL-Project.docx"
PDF   = PAPER / "URL-Project.pdf"

TITLE = "Machine Learning-Based Detection of Phishing URLs Using Lexical URL Features"

AUTHOR1 = [
    "Saleh Alsaheb",
    "Networks and Information Security Engineering",
    "Princess Sumaya University for Technology (PSUT)",
    "Amman, Jordan",
]
AUTHOR2 = [
    "Ismael Alhindi",
    "Networks and Information Security Engineering",
    "Princess Sumaya University for Technology (PSUT)",
    "Amman, Jordan",
]

ABSTRACT = (
    "Phishing web pages remain a practical threat because they imitate legitimate sites and "
    "attempt to steal credentials, financial information, or other sensitive data. This project "
    "evaluates whether lightweight URL-based lexical features, augmented with six engineered "
    "ratio and binary indicators, can support reliable phishing detection with classical "
    "supervised machine learning methods. The experiments use the Kaggle Web Page Phishing "
    "Dataset, which contains 100,077 rows, 19 URL-derived numerical features, and one binary "
    "target label. After removing 78,186 duplicate rows, the modeling workflow is run on 21,891 "
    "unique records. Six additional features are derived from the original columns to capture "
    "symbol density, redirection presence, and URL depth. The pipeline includes dataset "
    "inspection, duplicate handling, feature engineering, stratified train/test splitting, "
    "feature scaling where required, five-fold cross-validation, hyperparameter tuning via "
    "GridSearchCV for all four models, and held-out test evaluation. The four evaluated models "
    "are Logistic Regression, Decision Tree, Random Forest, and Gradient Boosting. Gradient "
    "Boosting achieved the strongest held-out test performance with 0.8198 accuracy, 0.8505 "
    "precision, 0.9118 recall, 0.8801 F1-score, and 0.8873 ROC-AUC. All four models achieved "
    "F1-scores between 0.853 and 0.880, demonstrating that the combination of feature "
    "engineering and systematic hyperparameter tuning produces consistent and competitive results "
    "across the evaluated model set."
)

BODY = [
    # I. Introduction
    ("H1", "I. Introduction"),
    ("P", "Phishing attacks are a common cybersecurity problem because they trick users into "
          "interacting with fraudulent pages that appear trustworthy. Attackers often rely on "
          "deceptive URLs, domain patterns, and redirection behavior to mislead victims. Because "
          "phishing campaigns can change quickly and operate at scale, automated detection is an "
          "important defensive layer."),
    ("P", "URL-based detection is attractive because it is lightweight and does not require "
          "downloading webpage content, executing scripts, or querying richer external sources "
          "such as WHOIS or DNS. Lexical properties such as URL length, the number of dots, the "
          "number of hyphens, and the presence of special characters can provide useful signals "
          "about suspicious behavior. These properties are also straightforward to extract and "
          "evaluate in a batch machine learning workflow."),
    ("P", "This project frames phishing detection as a supervised binary classification problem. "
          "The task is to classify a URL as legitimate or phishing, training and evaluation use a "
          "labeled public dataset, and the performance is measured using accuracy, precision, "
          "recall, F1-score, ROC-AUC, and confusion matrix analysis."),
    ("P", "Because a missed phishing page is more dangerous than a false alarm, "
          "recall is especially important."),
    ("P", "The main contributions of this project are:"),
    ("ENUM", [
        "A reproducible Python implementation for phishing URL detection using lexical URL features.",
        "A comparison of four classical and ensemble machine learning models under a shared evaluation workflow.",
        "A structured analysis of class balance, feature relationships, confusion matrix behavior, and recall/F1 trade-offs.",
        "A literature-based comparison with related phishing URL detection studies that use similar lexical features.",
    ]),

    # II. Related Work
    ("H1", "II. Related Work"),
    ("P", "Blum et al. studied phishing URL detection using lexical features and online "
          "learning methods [2]. Their work demonstrated that useful phishing signals can be extracted "
          "directly from URL text without relying on full webpage content. The present project is "
          "related in its emphasis on lexical features, but differs in dataset choice and in the "
          "broader comparison of batch classifiers."),
    ("P", "Sahingoz et al. examined machine learning based phishing detection from URLs "
          "using lexical and host-based signals [3]. In their reported experiments, Random Forest with "
          "NLP-based features achieved 97.98% accuracy on a large URL dataset containing 36,400 "
          "legitimate and 37,175 phishing URLs. The present project is narrower because it focuses "
          "on the Kaggle Web Page Phishing Dataset and a simple reproducible pipeline centered on "
          "URL lexical features."),
    ("P", "Awasthi and Goel evaluated both base and ensemble classifiers on two "
          "URL-feature phishing datasets using 30 structural URL features and 10-fold "
          "cross-validation [4]. Their base classifiers, including Logistic Regression, Decision "
          "Tree, and SVM, achieved accuracy in the range of 93 to 96 percent, which is "
          "methodologically relevant because it uses comparable classifier families and URL-based "
          "features, although the datasets and feature sets differ."),
    ("P", "Alnemari and Alshammari evaluated Random Forest, Decision Tree, SVM, and ANN "
          "on the UCI phishing website dataset of 11,055 instances described by 30 URL-structural "
          "features [5]. Random Forest achieved the highest result at 97.3 percent accuracy, while "
          "SVM and Decision Tree fell in the 94 to 96 percent range. Their work is relevant "
          "because it uses the same classical model families as the current project on a "
          "comparable URL-feature set."),
    ("P", "Although this work uses the Kaggle Web Page Phishing Dataset, the compared studies "
          "use different but related phishing URL datasets. The comparison therefore focuses on "
          "methodology, feature type, classifiers, and reported metric ranges rather than "
          "claiming a direct same-dataset benchmark."),

    # III. Experimental Setup
    ("H1", "III. Experimental Setup"),
    ("H2", "A. Dataset Description"),
    ("P", "The dataset used in this project is the Kaggle Web Page Phishing Dataset [1]. In the "
          "local CSV used for implementation, the dataset contains 100,077 rows and 20 columns "
          "before cleaning. Nineteen columns are numerical URL-derived features, and the target "
          "column is `phishing`, where 1 denotes phishing and 0 denotes legitimate. After "
          "duplicate removal, 21,891 unique rows remain for modeling. Six additional features "
          "are then engineered from the original columns, producing 25 features in total."),
    ("P", "Table I lists the input features and the target variable used in the pipeline."),
    ("TABLE", ([
        ["Feature", "Type", "Description"],
        ["url_length",          "Numerical",  "Total characters in the URL"],
        ["n_dots",              "Numerical",  "Count of . characters"],
        ["n_hypens",            "Numerical",  "Count of hyphen characters"],
        ["n_underline",         "Numerical",  "Count of _ characters"],
        ["n_slash",             "Numerical",  "Count of / characters"],
        ["n_questionmark",      "Numerical",  "Count of ? characters"],
        ["n_equal",             "Numerical",  "Count of = characters"],
        ["n_at",                "Numerical",  "Count of @ characters"],
        ["n_and",               "Numerical",  "Count of & characters"],
        ["n_exclamation",       "Numerical",  "Count of ! characters"],
        ["n_space",             "Numerical",  "Count of space characters"],
        ["n_tilde",             "Numerical",  "Count of ~ characters"],
        ["n_comma",             "Numerical",  "Count of , characters"],
        ["n_plus",              "Numerical",  "Count of + characters"],
        ["n_asterisk",          "Numerical",  "Count of * characters"],
        ["n_hastag",            "Numerical",  "Count of # characters"],
        ["n_dollar",            "Numerical",  "Count of $ characters"],
        ["n_percent",           "Numerical",  "Count of % characters"],
        ["n_redirection",       "Numerical",  "Count of redirect-like markers"],
        ["total_special_chars", "Engineered", "Sum of all special character counts"],
        ["symbol_ratio",        "Engineered", "total_special_chars / (url_length+1)"],
        ["has_redirect",        "Engineered", "Binary: 1 if n_redirection > 0"],
        ["has_at",              "Engineered", "Binary: 1 if n_at > 0"],
        ["slash_dot_ratio",     "Engineered", "n_slash / (n_dots+1)"],
        ["url_length_log",      "Engineered", "log(url_length + 1)"],
        ["phishing",            "Binary",     "1 = phishing, 0 = legitimate (target)"],
    ], [0.31, 0.15, 0.54], "FEATURES USED IN THE PHISHING DETECTION PIPELINE")),
    ("P", "No missing values were observed in the cleaned dataset. After duplicate removal, the "
          "class distribution was 6,019 legitimate URLs and 15,872 phishing URLs, corresponding "
          "to roughly 27.5% legitimate and 72.5% phishing samples."),
    ("FIG", ("class_distribution.png",
             "Class distribution after deduplication (6,019 legitimate vs 15,872 phishing).")),

    ("H2", "B. Exploratory Data Analysis"),
    ("P", "The exploratory analysis focuses on class balance, feature distributions, and "
          "feature-target relationships. The correlation analysis shows that `n_slash`, "
          "`url_length`, `n_equal`, `n_redirection`, and `n_dots` have the strongest absolute "
          "linear relationships with the phishing label in this dataset. Several count-based "
          "features exhibit noticeable skew and long tails, especially `url_length`, `n_hypens`, "
          "and `n_slash`."),
    ("FIG", ("top_correlations.png",
             "Absolute Pearson correlation of features with the phishing label (ranked).")),
    ("FIG", ("feature_distributions.png",
             "Feature distributions clipped to the 1st-99th percentile range.",
             1.80)),
    ("FIG", ("correlation_heatmap.png",
             "Feature correlation heatmap. Warmer colors (red) indicate positive correlation; cooler (blue) indicate negative.",
             1.10)),

    ("H2", "C. Data Preprocessing"),
    ("P", "Duplicate rows are removed once before model splitting, which is especially important "
          "in this dataset because more than three quarters of the original rows are duplicates. "
          "After deduplication, six features are engineered: a total special character count, a "
          "symbol-to-length ratio, binary redirect and at-symbol flags, a slash-to-dot ratio, "
          "and a log-transformed URL length. A stratified 80/20 train/test split is used to "
          "preserve class proportions. Feature scaling is applied only to Logistic Regression."),

    ("H2", "D. Evaluation Metrics"),
    ("P", "The project uses accuracy, precision, recall, F1-score, ROC-AUC, and confusion matrix "
          "analysis. Recall deserves special attention because false negatives correspond to "
          "phishing pages that evade detection, which carries the highest security risk."),

    # IV. Machine Learning Algorithms
    ("H1", "IV. Machine Learning Algorithms"),
    ("P", "Four classifiers are evaluated to cover the spectrum from simple linear models to "
          "complex boosting ensembles. Each model is tuned independently with GridSearchCV "
          "using 5-fold stratified cross-validation scored on F1."),

    ("H2", "A. Logistic Regression"),
    ("P", "Logistic Regression is included as the linear baseline. It models the log-odds of "
          "the phishing class as a weighted sum of the input features. Standard scaling is "
          "applied upstream in the pipeline. The regularization parameter C is tuned across "
          "0.01, 0.1, 1, and 10."),

    ("H2", "B. Decision Tree"),
    ("P", "The Decision Tree classifier recursively partitions the feature space by selecting "
          "the split that maximizes information gain at each node. Maximum depth, minimum "
          "samples per split, and minimum samples per leaf are all tuned to control overfitting."),

    ("H2", "C. Random Forest"),
    ("P", "Random Forest is a bagging ensemble that trains a collection of decision trees on "
          "bootstrapped subsets of the training data with randomized feature subsets at each "
          "split. The number of estimators, maximum depth, and leaf size parameters are tuned "
          "via grid search."),

    ("H2", "D. Gradient Boosting"),
    ("P", "Gradient Boosting is a sequential boosting ensemble that builds trees one at a time, "
          "each correcting the residual errors of the previous ensemble. The learning rate, "
          "number of estimators, maximum tree depth, and subsampling ratio are tuned together "
          "to balance bias and variance."),

    # V. Results and Analysis
    ("H1", "V. Results and Analysis"),
    ("H2", "A. Model Performance"),
    ("P", "Table II summarizes the held-out test performance of all four models after "
          "hyperparameter tuning via GridSearchCV."),
    ("TABLE", ([
        ["Model", "Accuracy", "Precision", "Recall", "F1", "AUC"],
        ["Gradient Boosting", "0.8198", "0.8505", "0.9118", "0.8801", "0.8873"],
        ["Random Forest",     "0.8102", "0.8429", "0.9074", "0.8740", "0.8793"],
        ["Decision Tree",     "0.8022", "0.8476", "0.8866", "0.8667", "0.8478"],
        ["Logistic Regression","0.7753","0.8101", "0.9014", "0.8533", "0.8329"],
    ], [0.35, 0.13, 0.13, 0.13, 0.12, 0.14], "TEST PERFORMANCE AFTER HYPERPARAMETER TUNING")),
    ("P", "Gradient Boosting achieved the best overall result, delivering the highest accuracy, "
          "F1-score, and ROC-AUC. Random Forest was a close second with F1 = 0.874. Decision "
          "Tree improved substantially from an untuned F1 of 0.791 to 0.867 after hyperparameter tuning. "
          "The overall F1 range of 0.853 to 0.880 indicates consistent and competitive "
          "performance across the evaluated model set."),
    ("P", "Since phishing samples form about 72.5% of the deduplicated dataset, accuracy alone "
          "is not sufficient; therefore, F1-score, recall, ROC-AUC, and confusion matrix "
          "behavior are emphasized throughout the evaluation."),
    ("FIG", ("model_comparison_accuracy.png",
             "Accuracy comparison across all four evaluated models.")),
    ("FIG", ("model_comparison_f1.png",
             "F1-score comparison across all four evaluated models.")),
    ("FIG", ("roc_curves.png",
             "ROC curves for all evaluated models. AUC values shown in legend.")),

    ("H2", "B. Best Model Confusion Matrix"),
    ("P", "Table III shows the confusion matrix for Gradient Boosting on the held-out test set."),
    ("TABLE", ([
        ["Actual / Predicted", "Legitimate", "Phishing"],
        ["Legitimate",         "695",        "509"],
        ["Phishing",           "280",        "2895"],
    ], [0.45, 0.275, 0.275], "CONFUSION MATRIX OF THE GRADIENT BOOSTING MODEL")),
    ("P", "The model correctly identifies 2,895 of 3,175 phishing URLs, yielding recall = 0.912. "
          "The 509 false positives represent legitimate URLs flagged as phishing, a "
          "security-oriented trade-off, although the false-positive rate should be reduced "
          "before real-world deployment."),
    ("FIG", ("confusion_matrix_best_model.png",
             "Confusion matrix for Gradient Boosting on the held-out test set.")),
    ("FIG", ("feature_importance.png",
             "Top feature importances for Gradient Boosting (best model).")),

    ("H2", "C. Discussion"),
    ("ENUM", [
        "Gradient Boosting achieved F1 = 0.8801 and ROC-AUC = 0.8873, the best results overall.",
        "Decision Tree improved substantially (F1: 0.791 to 0.867), demonstrating the importance of hyperparameter tuning.",
        "All four models are competitive with F1 spanning 0.853-0.880.",
        "Engineered features symbol_ratio and slash_dot_ratio appear in the top feature importances of tree-based models (Fig. 9).",
        "Results differ from [4] and [5] in absolute accuracy due to stricter deduplication and a more limited feature set, but the evaluated model families are methodologically comparable.",
    ]),

    # VI. Discussion and Limitations
    ("H1", "VI. Discussion and Limitations"),
    ("P", "This project uses URL lexical features augmented with six engineered indicators, "
          "which makes the method lightweight and fast. The engineered features, specifically "
          "the symbol density ratio, slash-to-dot ratio, and binary redirect and at-symbol "
          "flags, provide additional discriminative signals that complement the raw counts and "
          "benefit linear models that cannot learn such relationships on their own."),
    ("P", "Even with feature engineering, lexical features alone may miss sophisticated "
          "phishing pages that use cleaner-looking URLs or rely on infrastructure-level "
          "deception. The compared studies that report higher accuracy generally use richer "
          "feature sets that include structural, host-based, or content-based attributes. The "
          "current pipeline intentionally limits itself to URL-text features for "
          "reproducibility and computational simplicity."),
    ("P", "Another limitation is dataset freshness. A model trained on a static public dataset "
          "may not fully reflect the newest phishing strategies. The current pipeline does not "
          "inspect webpage content, WHOIS data, DNS signals, SSL information, or dynamic "
          "behavior. Real-world deployment would require external validation on more recent data."),
    ("P", "The duplicate structure of the provided CSV is a notable limitation. Removing "
          "duplicates reduced the working dataset from 100,077 rows to 21,891 unique rows. "
          "Results from experiments that do not deduplicate may appear artificially stronger."),

    # VII. Conclusion
    ("H1", "VII. Conclusion"),
    ("P", "This project demonstrates a complete machine learning workflow for phishing URL "
          "detection using a public dataset of lexical URL features augmented with six "
          "engineered indicators. After deduplication, experiments were conducted on 21,891 "
          "unique URL records. All four models were tuned via GridSearchCV, and hyperparameter "
          "optimization produced a substantial improvement in Decision Tree performance "
          "(F1: 0.791 to 0.867) and modest but consistent gains across the ensemble models."),
    ("P", "Among the evaluated models, Gradient Boosting achieved the strongest overall test "
          "performance with 0.8198 accuracy, 0.8505 precision, 0.9118 recall, 0.8801 "
          "F1-score, and 0.8873 ROC-AUC. These results are methodologically aligned with "
          "previous studies [4], [5], but the absolute accuracy is lower, likely because "
          "this project applies strict duplicate removal and uses a limited lexical feature set."),
    ("P", "URL lexical features combined with derived indicators can provide effective phishing "
          "detection without content-based inspection. Future work should validate the approach "
          "on newer datasets and expand the feature set with host, DNS, SSL, or content-based signals."),

    # References
    ("H1", "References"),
    ("REF", [
        '[1] Danielfernandon, "Web Page Phishing Dataset," Kaggle, 2023. [Online]. '
        'Available: https://www.kaggle.com/datasets/danielfernandon/web-page-phishing-dataset. '
        'Accessed: May 2026.',
        '[2] A. Blum, B. Wardman, T. Solorio, and G. Warner, "Lexical feature based phishing URL '
        'detection using online learning," in Proc. 3rd ACM Workshop on Artificial Intelligence '
        'and Security, 2010, pp. 54–60.',
        '[3] O. K. Sahingoz, E. Buber, O. Demir, and B. Diri, "Machine learning based phishing '
        'detection from URLs," Expert Systems with Applications, vol. 117, pp. 345–357, 2019.',
        '[4] A. Awasthi and N. Goel, "Phishing website prediction using base and ensemble classifier '
        'techniques with cross-validation," Cybersecurity, vol. 5, no. 1, p. 22, 2022.',
        '[5] S. Alnemari and M. Alshammari, "Detecting phishing domains using machine learning," '
        'Applied Sciences, vol. 13, no. 8, p. 4649, 2023.',
    ]),
]

_ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V"}


def _to_roman(n: int) -> str:
    return _ROMAN.get(n, str(n))


# ─── DOCX (converted from PDF) ───────────────────────────────────────────────

def make_docx(pdf_path: Path, out_path: Path) -> None:
    from pdf2docx import Converter
    cv = Converter(str(pdf_path))
    cv.convert(str(out_path), start=0, end=None)
    cv.close()
    print(f"Saved DOCX: {out_path}")


def _make_docx_legacy(out_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    FONT = "Times New Roman"

    doc = Document()
    sec = doc.sections[0]
    sec.page_width  = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin   = Inches(0.625)
    sec.right_margin  = Inches(0.625)
    sec.top_margin    = Inches(0.75)
    sec.bottom_margin = Inches(1.0)

    def _set_cols(section, n, space=360):
        sp = section._sectPr
        for c in sp.findall(qn('w:cols')):
            sp.remove(c)
        el = OxmlElement('w:cols')
        el.set(qn('w:num'), str(n))
        el.set(qn('w:space'), str(space))
        sp.append(el)

    _set_cols(sec, 2, 360)

    def _font(run, size=10, bold=False, italic=False):
        run.font.name   = FONT
        run.font.size   = Pt(size)
        run.font.bold   = bold
        run.font.italic = italic

    def _inline(para, text, size=10, bold=False, italic=False):
        for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text):
            if part.startswith("**") and part.endswith("**"):
                r = para.add_run(part[2:-2]); _font(r, size, bold=True, italic=italic)
            elif part.startswith("*") and part.endswith("*"):
                r = para.add_run(part[1:-1]); _font(r, size, bold=bold, italic=True)
            elif part.startswith("`") and part.endswith("`"):
                r = para.add_run(part[1:-1])
                r.font.name = "Courier New"; r.font.size = Pt(size - 1)
            elif part:
                r = para.add_run(part); _font(r, size, bold=bold, italic=italic)

    def _para(text="", align=WD_ALIGN_PARAGRAPH.LEFT, size=10,
               bold=False, italic=False, sb=0, sa=5):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(sb)
        p.paragraph_format.space_after  = Pt(sa)
        if text:
            _inline(p, text, size, bold, italic)
        return p

    def _hrule():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bot  = OxmlElement('w:bottom')
        bot.set(qn('w:val'),   'single')
        bot.set(qn('w:sz'),    '4')
        bot.set(qn('w:space'), '1')
        bot.set(qn('w:color'), '000000')
        pBdr.append(bot)
        pPr.append(pBdr)

    def _section_break_1col():
        """Continuous section break; declares preceding content as 1-column."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        pPr = p._p.get_or_add_pPr()
        sp  = OxmlElement('w:sectPr')
        t   = OxmlElement('w:type');   t.set(qn('w:val'), 'continuous')
        c   = OxmlElement('w:cols');   c.set(qn('w:num'), '1'); c.set(qn('w:space'), '360')
        pSz = OxmlElement('w:pgSz');   pSz.set(qn('w:w'), '12240'); pSz.set(qn('w:h'), '15840')
        pMar = OxmlElement('w:pgMar')
        for attr, val in [('w:top','1080'),('w:right','900'),('w:bottom','1440'),('w:left','900')]:
            pMar.set(qn(attr), val)
        for el in [t, c, pSz, pMar]:
            sp.append(el)
        pPr.append(sp)

    fig_n = [0]

    def _figure(filename, caption, *_):
        fig_path = FIGS / filename
        if not fig_path.exists():
            return
        fig_n[0] += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(2)
        try:
            p.add_run().add_picture(str(fig_path), width=Inches(2.85))
        except Exception:
            p.add_run(f"[Figure {fig_n[0]}: {filename}]")
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        r = cap.add_run(f"Fig. {fig_n[0]}. {caption}")
        _font(r, 8, italic=True)

    tbl_n = [0]

    def _table(data):
        tbl_title = None
        if isinstance(data, tuple) and len(data) == 3:
            rows, col_ratios, tbl_title = data
        elif isinstance(data, tuple) and len(data) == 2 and isinstance(data[1], list):
            rows, col_ratios = data
        else:
            rows, col_ratios = data, None
        if not rows:
            return
        tbl_n[0] += 1
        _para(f"TABLE {_to_roman(tbl_n[0])}", align=WD_ALIGN_PARAGRAPH.CENTER, size=8, bold=True, sb=6, sa=1)
        if tbl_title:
            _para(tbl_title, align=WD_ALIGN_PARAGRAPH.CENTER, size=8, bold=True, sb=0, sa=2)
        DOC_COL_W = 3.5
        n_cols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=n_cols)
        t.style = 'Table Grid'
        for ri, row_data in enumerate(rows):
            for ci, cell_text in enumerate(row_data[:n_cols]):
                cell = t.rows[ri].cells[ci]
                cell.text = ""
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(1)
                r = cp.add_run(cell_text)
                _font(r, 8, bold=(ri == 0))
                if col_ratios and ci < len(col_ratios):
                    tcPr = cell._tc.get_or_add_tcPr()
                    for w in tcPr.findall(qn('w:tcW')):
                        tcPr.remove(w)
                    tcW = OxmlElement('w:tcW')
                    tcW.set(qn('w:w'), str(int(col_ratios[ci] * DOC_COL_W * 1440)))
                    tcW.set(qn('w:type'), 'dxa')
                    tcPr.append(tcW)
                if ri == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd  = OxmlElement('w:shd')
                    shd.set(qn('w:val'),   'clear')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:fill'),  'CCCCCC')
                    tcPr.append(shd)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ── Title block (single-column) ───────────────────────────────────────────
    _para(TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, size=18, bold=True, sb=0, sa=6)

    # Two-author block side by side (borderless 2-col table)
    auth_tbl = doc.add_table(rows=1, cols=2)
    auth_tbl.style = 'Table Grid'
    for ci, author_lines in enumerate([AUTHOR1, AUTHOR2]):
        cell = auth_tbl.rows[0].cells[ci]
        cell.text = ""
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement(f'w:{side}')
            el.set(qn('w:val'), 'none')
            tcBorders.append(el)
        tcPr.append(tcBorders)
        for li, line in enumerate(author_lines):
            if li == 0:
                cp = cell.paragraphs[0]
            else:
                cp = cell.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(1)
            r = cp.add_run(line)
            _font(r, 10, bold=(li == 0), italic=(li > 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    abs_p = doc.add_paragraph()
    abs_p.paragraph_format.space_after   = Pt(4)
    abs_p.paragraph_format.left_indent   = Inches(0.5)
    abs_p.paragraph_format.right_indent  = Inches(0.5)
    r1 = abs_p.add_run("Abstract—"); _font(r1, 9, bold=True)
    r2 = abs_p.add_run(ABSTRACT);        _font(r2, 9, italic=True)

    _hrule()
    _section_break_1col()

    # ── Body ─────────────────────────────────────────────────────────────────
    for kind, data in BODY:
        if kind == "H1":
            _para(data.upper(), align=WD_ALIGN_PARAGRAPH.CENTER, size=10, bold=True, sb=10, sa=4)
        elif kind == "H2":
            _para(data, size=10, bold=True, italic=True, sb=6, sa=3)
        elif kind == "P":
            _para(data, size=10, sb=0, sa=5)
        elif kind == "ENUM":
            for item in data:
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(3)
                _inline(p, item, 10)
        elif kind == "TABLE":
            _table(data)
        elif kind == "FIG":
            _figure(*data)
        elif kind == "REF":
            for item in data:
                p = _para(item, size=9, sa=2)
                p.paragraph_format.left_indent       = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)

    doc.save(str(out_path))
    print(f"Saved DOCX: {out_path}")


# ─── PDF ─────────────────────────────────────────────────────────────────────

def make_pdf(out_path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, Table,
        TableStyle, Image, KeepTogether, HRFlowable, FrameBreak, NextPageTemplate,
    )
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
    from reportlab.lib.colors import HexColor, black, white

    W, H = letter
    LM = RM = 0.625 * inch
    TM = 0.75 * inch
    BM = 1.0  * inch
    GAP = 0.25 * inch

    BW   = W - LM - RM
    COLW = (BW - GAP) / 2
    FULL_H = H - TM - BM
    HDR_H  = 2.15 * inch        # header frame for title+authors on p.1 (abstract in col 1)
    COL_H1 = FULL_H - HDR_H    # column height on p.1 below header

    hdr   = Frame(LM, H - TM - HDR_H, BW,   HDR_H,  id='hdr',  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    lp1   = Frame(LM, BM,              COLW,  COL_H1, id='lp1',  leftPadding=0, rightPadding=3, topPadding=0, bottomPadding=0)
    rp1   = Frame(LM + COLW + GAP, BM, COLW,  COL_H1, id='rp1',  leftPadding=3, rightPadding=0, topPadding=0, bottomPadding=0)
    lft   = Frame(LM, BM,              COLW,  FULL_H,  id='lft',  leftPadding=0, rightPadding=3, topPadding=0, bottomPadding=0)
    rgt   = Frame(LM + COLW + GAP, BM, COLW,  FULL_H,  id='rgt',  leftPadding=3, rightPadding=0, topPadding=0, bottomPadding=0)

    def _on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Times-Roman", 9)
        canvas.drawCentredString(W / 2, BM / 2, str(doc.page))
        canvas.restoreState()

    doc = BaseDocTemplate(str(out_path), pagesize=letter,
                          leftMargin=LM, rightMargin=RM, topMargin=TM, bottomMargin=BM)
    doc.addPageTemplates([
        PageTemplate(id='first', frames=[hdr, lp1, rp1], onPage=_on_page),
        PageTemplate(id='body',  frames=[lft, rgt],       onPage=_on_page),
    ])

    def S(**kw):
        kw.setdefault('fontName', 'Times-Roman')
        kw.setdefault('fontSize', 10)
        kw.setdefault('leading',  14)
        kw.setdefault('spaceAfter', 5)
        kw.setdefault('alignment', TA_JUSTIFY)
        return ParagraphStyle(kw.pop('name', '_'), **kw)

    sTitle = S(name='T',  fontName='Times-Bold', fontSize=18, leading=22,
                alignment=TA_CENTER, spaceAfter=6)
    sAuth  = S(name='A',  fontName='Times-Bold', fontSize=11, leading=14,
                alignment=TA_CENTER, spaceAfter=1)
    sAuthSub = S(name='As', fontName='Times-Italic', fontSize=10, leading=13,
                alignment=TA_CENTER, spaceAfter=1)
    sAbs   = S(name='Ab', fontName='Times-Italic', fontSize=9, leading=13,
                alignment=TA_JUSTIFY, spaceAfter=4)
    sH1    = S(name='H1', fontName='Times-Bold', fontSize=10, leading=14,
                alignment=TA_CENTER, spaceBefore=6, spaceAfter=4, keepWithNext=1)
    sH2    = S(name='H2', fontName='Times-BoldItalic', fontSize=10, leading=14,
                alignment=TA_LEFT, spaceBefore=4, spaceAfter=3, keepWithNext=1)
    sBody  = S(name='Bd', fontSize=10, leading=14, spaceAfter=5)
    sCap   = S(name='Cp', fontName='Times-Italic', fontSize=8, leading=11,
                alignment=TA_CENTER, spaceAfter=8)
    sRef   = S(name='Rf', fontSize=9, leading=12, spaceAfter=2,
                leftIndent=14, firstLineIndent=-14)
    sTbl   = S(name='Th', fontName='Times-Bold', fontSize=8, leading=11,
                alignment=TA_CENTER, spaceAfter=2)

    def _h(text: str) -> str:
        text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
        text = re.sub(r'\*(.+?)\*',     r'<i>\1</i>', text)
        text = re.sub(r'`(.+?)`', r'<font face="Courier" size="9">\1</font>', text)
        return text

    story = []

    # Header (fills hdr frame on page 1)
    story.append(Paragraph(TITLE, sTitle))
    story.append(Spacer(1, 6))

    # Two-author block side by side
    def _author_cell_paras(author_lines):
        paras = []
        for li, line in enumerate(author_lines):
            style = sAuth if li == 0 else sAuthSub
            paras.append(Paragraph(line, style))
        return paras

    auth_data = [
        [_author_cell_paras(AUTHOR1), _author_cell_paras(AUTHOR2)]
    ]
    auth_col_w = BW / 2
    auth_tbl = Table(auth_data, colWidths=[auth_col_w, auth_col_w])
    auth_tbl.setStyle(TableStyle([
        ('VALIGN',       (0, 0), (-1, -1), 'TOP'),
        ('ALIGN',        (0, 0), (-1, -1), 'CENTER'),
        ('LEFTPADDING',  (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING',   (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING',(0, 0), (-1, -1), 4),
        ('LINEBELOW',    (0, 0), (-1, -1), 0, white),
    ]))
    story.append(auth_tbl)
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=0.5, color=black))
    story.append(Spacer(1, 4))
    story.append(FrameBreak())          # done with header; move to lp1
    story.append(NextPageTemplate('body'))  # subsequent pages use two full columns
    # Abstract flows into col 1 of page 1
    story.append(Paragraph(f'<b>Abstract—</b><i>{ABSTRACT}</i>', sAbs))
    story.append(Spacer(1, 8))

    fig_n = [0]
    tbl_n = [0]

    for kind, data in BODY:
        if kind == "H1":
            story.append(Paragraph(data.upper(), sH1))
        elif kind == "H2":
            story.append(Paragraph(data, sH2))
        elif kind == "P":
            story.append(Paragraph(_h(data), sBody))
        elif kind == "ENUM":
            for i, item in enumerate(data, 1):
                story.append(Paragraph(f"{i}. {_h(item)}", sBody))
        elif kind == "TABLE":
            tbl_n[0] += 1
            tbl_title = None
            if isinstance(data, tuple) and len(data) == 3:
                rows, col_ratios, tbl_title = data
            elif isinstance(data, tuple) and len(data) == 2 and isinstance(data[1], list):
                rows, col_ratios = data
            else:
                rows, col_ratios = data, None
            ncols = max(len(r) for r in rows)
            if col_ratios:
                col_w = [COLW * r for r in col_ratios]
            else:
                col_w = [COLW / ncols] * ncols
            t = Table(rows, colWidths=col_w, repeatRows=1)
            t.setStyle(TableStyle([
                ('FONTNAME',      (0, 0), (-1,  0), 'Times-Bold'),
                ('FONTNAME',      (0, 1), (-1, -1), 'Times-Roman'),
                ('FONTSIZE',      (0, 0), (-1, -1), 7),
                ('LEADING',       (0, 0), (-1, -1), 10),
                ('BACKGROUND',    (0, 0), (-1,  0), HexColor('#CCCCCC')),
                ('GRID',          (0, 0), (-1, -1), 0.5, black),
                ('ALIGN',         (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING',    (0, 0), (-1, -1), 2),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
                ('LEFTPADDING',   (0, 0), (-1, -1), 3),
                ('RIGHTPADDING',  (0, 0), (-1, -1), 3),
            ]))
            tbl_items = [Paragraph(f"TABLE {_to_roman(tbl_n[0])}", sTbl)]
            if tbl_title:
                tbl_items.append(Paragraph(tbl_title, sTbl))
            tbl_items.extend([t, Spacer(1, 5)])
            story.append(KeepTogether(tbl_items))
        elif kind == "FIG":
            filename = data[0]
            caption  = data[1]
            h_factor = data[2] if len(data) > 2 else None
            fig_path = FIGS / filename
            if fig_path.exists():
                fig_n[0] += 1
                img = Image(str(fig_path))
                aspect = img.imageHeight / float(img.imageWidth)
                target_w = COLW * 0.9
                target_h = target_w * aspect
                if h_factor is not None:
                    max_h = COLW * h_factor
                else:
                    max_h = COLW * (0.95 if aspect > 1 else 0.55)
                if target_h > max_h:
                    target_h = max_h
                    target_w = target_h / aspect
                img.drawWidth  = target_w
                img.drawHeight = target_h
                story.append(KeepTogether([
                    Spacer(1, 4),
                    img,
                    Paragraph(f"Fig. {fig_n[0]}. {caption}", sCap),
                ]))
        elif kind == "REF":
            for item in data:
                story.append(Paragraph(_h(item), sRef))

    doc.build(story)
    print(f"Saved PDF:  {out_path}")


# ─── MAIN ────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    make_pdf(PDF)
    _make_docx_legacy(DOCX)
    print("Done.")
